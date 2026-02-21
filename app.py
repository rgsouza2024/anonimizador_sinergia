# Nome do arquivo: app.py (anonimizador_sinergia)
# Versão 0.96 - Versão Beta

import gradio as gr
import spacy
from presidio_analyzer import AnalyzerEngine, PatternRecognizer
from presidio_analyzer.nlp_engine import SpacyNlpEngine
from presidio_analyzer.pattern import Pattern
from presidio_anonymizer import AnonymizerEngine
from presidio_anonymizer.entities import OperatorConfig
import pandas as pd
import io
import fitz  # PyMuPDF
from docx import Document
import re
import os
import time
import unicodedata
from dotenv import load_dotenv

# Carrega as variáveis de ambiente do arquivo .env (opcional, mas bom manter)
load_dotenv()

# --- Constantes ---
NOME_ARQUIVO_SOBRENOMES = "sobrenomes_comuns.txt"
NOME_ARQUIVO_TERMOS_COMUNS = "termos_comuns.txt"
NOME_ARQUIVO_TITULOS = "titulos_legais.txt"
LOGO_FILE_PATH = "Logo Projeto Sinergia TRF1 - Fundo Transparente.png"
ESTADO_VAZIO_TEXTO_ANONIMIZADO = "O texto anonimizado aparecera aqui apos clicar em 'Anonimizar texto'."
ESTADO_VAZIO_PDF_ORIGINAL = "O texto extraido do PDF sera exibido aqui apos o processamento."
ESTADO_VAZIO_PDF_ANONIMIZADO = "O resultado anonimizado do PDF sera exibido aqui apos o processamento."
COLUNAS_ENTIDADES_DETECTADAS = ["Entidade", "Texto Detectado", "Inicio", "Fim", "Score"]
RESUMO_VAZIO_TEXTO = "**Resumo:** aguardando processamento do texto."
RESUMO_VAZIO_PDF = "**Resumo:** aguardando processamento do PDF."
PLACEHOLDER_NOME_PARTE_EXTERNO = "<NOME_PARTE_AUTORA>"
PLACEHOLDER_NOME_PARTE_INTERNO = "__NOME_PARTE_AUTORA__"
PAPEIS_ALVO_NOME_PARTE = [
    "AUTOR",
    "AUTORA",
    "PARTE AUTORA",
    "RECORRENTE",
    "REQUERENTE",
    "RECORRIDO",
    "RECORRIDA",
]
TERMOS_EXCLUSAO_NOME_PARTE = [
    "ADVOGADO",
    "ADVOGADA",
    "OAB",
    "JUIZ",
    "JUIZA",
    "RELATOR",
    "PROCURADOR",
    "PROCURADORIA",
    "DESEMBARGADOR",
    "MINISTRO",
    "DEFENSOR",
    "PATRONO",
    "PATRONA",
    "TERCEIRO INTERESSADO",
]
TERMOS_INDICADORES_PJ = [
    "LTDA",
    "S/A",
    "EIRELI",
    "ME",
    "MEI",
    "CNPJ",
    "INSTITUTO",
    "ASSOCIACAO",
    "SOCIEDADE",
    "EMPRESA",
    "FUNDACAO",
    "AUTARQUIA",
    "UNIAO",
    "MUNICIPIO",
    "ESTADO",
    "MINISTERIO",
    "SECRETARIA",
]
PALAVRAS_NAO_NOME_GENERICAS = [
    "PODER", "JUDICIARIO", "JUSTICA", "FEDERAL", "SECAO", "VARA", "PROCESSO",
    "BENEFICIO", "ASSISTENCIAL", "PRESTACAO", "CONTINUADA", "SENTENCA", "CONTESTACAO",
    "TURMA", "RECURSAL", "TRIBUNAL", "REGIAO", "DOCUMENTO", "ASSINADO", "ELETRONICAMENTE",
    "ADVOCACIA", "PROCURADORIA", "INSS", "MINISTERIO", "PUBLICO", "DEFENSORIA", "UNIAO",
]
REGEX_UF_BR = r"(?:AC|AL|AP|AM|BA|CE|DF|ES|GO|MA|MT|MS|MG|PA|PB|PR|PE|PI|RJ|RN|RS|RO|RR|SC|SP|SE|TO)"

# --- Funções e Listas para o Motor de Anonimização ---
def carregar_lista_de_arquivo(nome_arquivo):
    lista_itens = []
    caminho_base = os.path.dirname(__file__)
    caminho_arquivo = os.path.join(caminho_base, nome_arquivo)
    try:
        with open(caminho_arquivo, "r", encoding="utf-8") as f:
            for linha in f:
                item = linha.strip()
                if item:
                    lista_itens.append(item)
    except FileNotFoundError:
        print(f"AVISO: Arquivo de lista '{nome_arquivo}' não encontrado.")
    except Exception as e:
        print(f"ERRO: Erro ao ler o arquivo '{nome_arquivo}': {e}")
    return lista_itens


LISTA_SOBRENOMES_FREQUENTES_BR = carregar_lista_de_arquivo(NOME_ARQUIVO_SOBRENOMES)
LISTA_TERMOS_COMUNS = carregar_lista_de_arquivo(NOME_ARQUIVO_TERMOS_COMUNS)
LISTA_TITULOS_LEGAIS = carregar_lista_de_arquivo(NOME_ARQUIVO_TITULOS)
LISTA_ESTADOS_CAPITAIS_BR = ["Acre", "AC", "Alagoas", "AL", "Amapá", "AP", "Amazonas", "AM", "Bahia", "BA", "Ceará", "CE", "Distrito Federal", "DF", "Espírito Santo", "ES", "Goiás", "GO", "Maranhão", "MA", "Mato Grosso", "MT", "Mato Grosso do Sul", "MS", "Minas Gerais", "MG", "Pará", "PA", "Paraíba", "PB", "Paraná", "PR", "Pernambuco", "PE", "Piauí", "PI", "Rio de Janeiro", "RJ", "Rio Grande do Norte", "RN", "Rio Grande do Sul", "RS", "Rondônia", "RO", "Roraima", "RR", "Santa Catarina", "SC", "São Paulo", "SP", "Sergipe", "SE", "Tocantins", "TO",
                             "Aracaju", "Belém", "Belo Horizonte", "Boa Vista", "Brasília", "Campo Grande", "Cuiabá", "Curitiba", "Florianópolis", "Fortaleza", "Goiânia", "João Pessoa", "Macapá", "Maceió", "Manaus", "Natal", "Palmas", "Porto Alegre", "Porto Velho", "Recife", "Rio Branco", "Salvador", "São Luís", "Teresina", "Vitória"]
TERMOS_CABECALHO_LEGAL_NAO_ANONIMIZAR = ["EXMO. SR. DR. JUIZ FEDERAL", "EXMO SR DR JUIZ FEDERAL", "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ FEDERAL", "JUIZ FEDERAL", "EXMO. SR. DR. JUIZ DE DIREITO", "EXMO SR DR JUIZ DE DIREITO", "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO", "JUIZ DE DIREITO", "JUIZADO ESPECIAL FEDERAL", "VARA DA SEÇÃO JUDICIÁRIA", "SEÇÃO JUDICIÁRIA", "EXMO.", "EXMO", "SR.", "DR.", "Dra.", "DRA.", "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A) FEDERAL",
                                       "EXCELENTÍSSIMO", "Senhor", "Doutor", "Senhora", "Doutora", "EXCELENTÍSSIMA", "EXCELENTÍSSIMO(A)", "Senhor(a)", "Doutor(a)", "Juiz", "Juíza", "Juiz(a)", "Juiz(íza)", "Assunto", "Assuntos"]
LISTA_ESTADO_CIVIL = ["casado", "casada", "solteiro", "solteira", "viúvo", "viúva", "divorciado", "divorciada",
                      "separado", "separada", "unido", "unida", "companheiro", "companheira", "amasiado", "amasiada", "união estável", "em união estável"]
LISTA_ORGANIZACOES_CONHECIDAS = ["FUNASA", "INSS", "IBAMA", "CNPQ", "IBGE", "FIOCRUZ", "SERPRO", "DATAPREV", "VALOR", "Justiça", "Justica", "Segredo", "PJe", "Assunto", "Tribunal Regional Federal", "Assuntos", "Vara Federal", "Vara", "Justiça Federal", "Federal", "Juizado", "Especial", "Federal", "Vara Federal de Juizado Especial Cível", "Turma",
                                "Turma Recursal", "PJE", "SJGO", "SJDF", "SJMA", "SJAC", "SJAL", "SJAP", "SJAM", "SJBA", "SJCE", "SJDF", "SJES", "SJGO", "SJMA", "SJMG", "SJMS", "SJMT", "SJPA", "SJPB", "SJPE", "SJPI", "SJPR", "SJPE", "SJRN", "SJRO", "SJRR", "SJRS", "SJSC", "SJSE", "SJSP", "SJTO", "Justiça Federal da 1ª Região", "PJe - Processo Judicial Eletrônico"]

# --- Configuração e Inicialização do Presidio (Motor Principal) ---
def carregar_analyzer_engine(termos_safe_location, termos_legal_header, lista_sobrenomes, termos_estado_civil, termos_organizacoes_conhecidas, termos_comuns_a_manter, titulos_legais):
    try:
        spacy.load('pt_core_news_lg')
    except OSError:
        print("ERRO CRÍTICO: Modelo spaCy 'pt_core_news_lg' não encontrado. Instale com: python -m spacy download pt_core_news_lg")
        return None
    spacy_engine_obj = SpacyNlpEngine(
        models=[{'lang_code': 'pt', 'model_name': 'pt_core_news_lg'}])
    analyzer = AnalyzerEngine(
        nlp_engine=spacy_engine_obj, supported_languages=["pt"], default_score_threshold=0.4)
    if termos_safe_location:
        analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="SAFE_LOCATION", name="SafeLocationRecognizer",
                                                            deny_list=termos_safe_location, supported_language="pt", deny_list_score=0.99))
    if termos_legal_header:
        analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="LEGAL_HEADER", name="LegalHeaderRecognizer",
                                                            deny_list=termos_legal_header, supported_language="pt", deny_list_score=0.99))

    if titulos_legais:
        palavras_para_regex = '|'.join(titulos_legais)
        legal_titles_pattern = Pattern(
            name="legal_title_pattern",
            regex=rf"(?i)\b({palavras_para_regex})(?:\([A-Z]\))?\b",
            score=1.0
        )
        legal_title_recognizer = PatternRecognizer(
            supported_entity="LEGAL_TITLE",
            name="LegalTitleRecognizer",
            patterns=[legal_titles_pattern],
            supported_language="pt"
        )
        analyzer.registry.add_recognizer(legal_title_recognizer)

    analyzer.registry.add_recognizer(PatternRecognizer(
        supported_entity="CPF",
        name="CustomCpfRecognizer",
        patterns=[
            Pattern(
                name="cpf_formatado",
                regex=r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b",
                score=1.0
            ),
            Pattern(
                name="cpf_com_separadores",
                regex=r"\b\d{3}[.\s]\d{3}[.\s]\d{3}[-\s]\d{2}\b",
                score=0.99
            ),
            Pattern(
                name="cpf_com_rotulo",
                regex=r"(?i)\bCPF\s*(?:n[ºo°.]?\s*)?:?\s*\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b",
                score=1.0
            ),
            Pattern(
                name="cin_com_rotulo",
                regex=r"(?i)\bCIN\s*(?:n[ºo°.]?\s*)?:?\s*\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b",
                score=1.0
            ),
            Pattern(
                name="carteira_identidade_nacional_com_numero",
                regex=r"(?i)\bCarteira\s+de\s+Identidade\s+Nacional\b[\s:\-nºo°\.]{0,20}\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b",
                score=0.99
            ),
            Pattern(
                name="cpf_apenas_numeros",
                regex=r"(?<!\d)\d{11}(?!\d)",
                score=0.995
            ),
        ],
        supported_language="pt"
    ))
    analyzer.registry.add_recognizer(PatternRecognizer(
        supported_entity="OAB_NUMBER",
        name="CustomOabRecognizer",
        patterns=[
            Pattern(
                name="oab_uf_numero",
                regex=rf"(?i)\bOAB\s*(?:n[º°.]?\s*)?(?:/|-)?\s*{REGEX_UF_BR}\s*(?:[Nn][º°o.\?]?\s*)?[-.]?\s*\d{{1,6}}(?:\.\d{{3}})?\b",
                score=0.90
            ),
            Pattern(
                name="oab_numero_uf",
                regex=rf"(?i)\bOAB\s*(?:n[º°.]?\s*)?\d{{1,6}}(?:\.\d{{3}})?\s*\/\s*{REGEX_UF_BR}\b",
                score=0.88
            ),
            Pattern(
                name="oab_uf_numero_compacto",
                regex=rf"(?i)\bOAB\s*(?:/|-)?\s*{REGEX_UF_BR}\s*[-.]?\s*\d{{4,6}}(?:\.\d{{3}})?\b",
                score=0.80
            ),
            Pattern(
                name="oab_uf_numero_sem_prefixo",
                regex=rf"(?i)(?<!/)\b{REGEX_UF_BR}\s*[-.]?\s*\d{{4,6}}(?:\.\d{{3}})?\b",
                score=0.70
            ),
        ],
        supported_language="pt"
    ))
    analyzer.registry.add_recognizer(PatternRecognizer(
        supported_entity="CEP_NUMBER",
        name="CustomCepRecognizer",
        patterns=[
            Pattern(
                name="cep_com_rotulo",
                regex=r"(?i)\bCEP\s*[:\-]?\s*\d{5}-?\d{3}\b",
                score=0.99
            ),
            Pattern(
                name="cep_hifen",
                regex=r"\b\d{5}-\d{3}\b",
                score=0.93
            ),
            Pattern(
                name="cep_pontuado",
                regex=r"\b\d{2}\.\d{3}-\d{3}\b",
                score=0.93
            ),
        ],
        supported_language="pt"
    ))
    analyzer.registry.add_recognizer(PatternRecognizer(
        supported_entity="ENDERECO_LOGRADOURO",
        name="EnderecoLogradouroRecognizer",
        patterns=[
            Pattern(
                name="logradouro_com_numero_e_complemento",
                regex=r"(?i)\b(?:rua|ra|avenida|av\.?|travessa|trav\.?|alameda|pra[cç]a|rodovia|estrada)\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+){0,8}\s*,\s*(?:n[ºo°.]?\s*)?\d+[A-Za-z0-9/\-]*\s*,\s*[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]{2,}(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]{2,}){0,5}",
                score=0.97
            ),
            Pattern(
                name="logradouro_prefixado",
                regex=r"(?i)\b(?:rua|ra|avenida|av\.?|travessa|trav\.?|alameda|pra[cç]a|rodovia|estrada|vicinal)\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+){0,8}(?:\s*,?\s*(?:n[ºo°.]?\s*)?\d+[A-Za-z0-9/\-]*)?(?:\s*,\s*[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]{2,}(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]{2,}){0,5})?",
                score=0.90
            ),
            Pattern(
                name="vicinal_composta",
                regex=r"(?i)\bVicinal\s+\d{1,3}\s+com\s+Vicinal\s+\d{1,3}\b",
                score=0.95
            ),
            Pattern(
                name="quadra_lote",
                regex=r"(?i)\b(?:quadra|qd\.?)\s*[A-Za-z0-9\-]+(?:\s*,?\s*(?:lote|lt\.?)\s*[A-Za-z0-9\-]+)?\b",
                score=0.88
            ),
            Pattern(
                name="setor_condominio_residencial",
                regex=r"(?i)\b(?:setor|condom[ií]nio|residencial)\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ0-9'`´.\-]+){0,5}\b",
                score=0.82
            ),
        ],
        supported_language="pt"
    ))
    if termos_estado_civil:
        analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="ESTADO_CIVIL", name="EstadoCivilRecognizer", patterns=[
                                      Pattern(name=f"estado_civil_{t.lower()}", regex=rf"(?i)\b{re.escape(t)}\b", score=0.99) for t in termos_estado_civil], supported_language="pt"))
    if termos_organizacoes_conhecidas:
        analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="ORGANIZACAO_CONHECIDA", name="OrganizacaoConhecidaRecognizer", patterns=[
                                      Pattern(name=f"org_{t.lower()}", regex=rf"(?i)\b{re.escape(t)}\b", score=0.99) for t in termos_organizacoes_conhecidas], supported_language="pt"))
    if termos_comuns_a_manter:
        recognizer_common_terms = PatternRecognizer(
            supported_entity="TERMO_COMUM", name="CommonTermsRecognizer", deny_list=termos_comuns_a_manter, supported_language="pt", deny_list_score=1.0)
        analyzer.registry.add_recognizer(recognizer_common_terms)
    analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="CNH", name="CNHRecognizer", patterns=[
        Pattern(name="cnh_formatado", regex=r"\bCNH\s*(?:nº|n\.)?\s*\d{11}\b", score=0.99),
    ], supported_language="pt"))
    analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="SIAPE", name="SIAPERecognizer", patterns=[Pattern(
        name="siape_formatado", regex=r"\bSIAPE\s*(?:nº|n\.)?\s*\d{7}\b", score=0.98), Pattern(name="siape_apenas_numeros", regex=r"\b(?<![\w])\d{7}(?![\w])\b", score=0.85)], supported_language="pt"))
    analyzer.registry.add_recognizer(PatternRecognizer(
        supported_entity="RG_NUMBER",
        name="CustomRgRecognizer",
        patterns=[
            Pattern(
                name="rg_com_rotulo",
                regex=r"(?i)\bRG\s*(?:n[ºo°.]?\s*)?[:\-]?\s*(?=[A-Z0-9.\-X]{3,20}\b)(?=[A-Z0-9.\-X]*\d)[A-Z0-9.\-X]{3,20}(?:\s*-\s*\dª\s*VIA)?(?:\s*[-,]?\s*(?:SSP|SESP|PC|IFP|DGPC|SJS|SJTC|SSDS|POL[ÍI]CIA\s+CIVIL)\s*\/\s*[A-Z]{2})?\b",
                score=0.99
            ),
            Pattern(
                name="registro_geral_com_numero",
                regex=r"(?i)\bRegistro\s+Geral\s*(?:n[ºo°.]?\s*)?[:\-]?\s*(?=[A-Z0-9.\-X]{3,20}\b)(?=[A-Z0-9.\-X]*\d)[A-Z0-9.\-X]{3,20}(?:\s*[-,]?\s*(?:SSP|SESP|PC|IFP|DGPC|SJS|SJTC|SSDS|POL[ÍI]CIA\s+CIVIL)\s*\/\s*[A-Z]{2})?\b",
                score=0.98
            ),
            Pattern(
                name="cedula_ou_carteira_identidade_com_numero",
                regex=r"(?i)\b(?:C[ÉE]DULA\s+DE\s+IDENTIDADE|CARTEIRA\s+DE\s+IDENTIDADE(?!\s+NACIONAL))\s*(?:n[ºo°.]?\s*)?[:\-]?\s*(?=[A-Z0-9.\-X]{3,20}\b)(?=[A-Z0-9.\-X]*\d)[A-Z0-9.\-X]{3,20}(?:\s*[-,]?\s*(?:SSP|SESP|PC|IFP|DGPC|SJS|SJTC|SSDS|POL[ÍI]CIA\s+CIVIL)\s*\/\s*[A-Z]{2})?\b",
                score=0.95
            )
        ],
        supported_language="pt"
    ))
    analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="MATRICULA_SIAPE", name="MatriculaSiapeRecognizer", patterns=[
                                      Pattern(name="matricula_siape", regex=r"(?i)\b(matr[íi]cula|siape)\b", score=0.95)], supported_language="pt"))
    analyzer.registry.add_recognizer(PatternRecognizer(supported_entity="ID_DOCUMENTO", name="IdDocumentoRecognizer", patterns=[
        Pattern(name="numero_beneficio_nb_formatado", regex=r"\bNB\s*\d{1,3}(\.?\d{3}){2}-[\dX]\b", score=0.98),
        Pattern(name="id_numerico_longo_pje", regex=r"\b\d{10,25}\b", score=0.97),
        Pattern(name="id_prefixo_numerico", regex=r"\bID\s*\d{8,12}\b", score=0.97),
        Pattern(name="numero_processo_cnj", regex=r"\b\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}\b", score=0.95),
        Pattern(name="numero_rnm", regex=r"\bRNM\s*(?:nº|n\.)?\s*[A-Z0-9]{7,15}\b", score=0.98),
        Pattern(name="numero_crm", regex=r"\bCRM\s*[A-Z]{2}\s*-\s*\d{1,6}\b", score=0.98)
    ], supported_language="pt"))
    return analyzer


def obter_operadores_anonimizacao():
    return {
        "DEFAULT": OperatorConfig("keep"),
        "PERSON": OperatorConfig("keep"),
        "LOCATION": OperatorConfig("keep"),
        "ENDERECO_LOGRADOURO": OperatorConfig("replace", {"new_value": "<ENDERECO>"}),
        "EMAIL_ADDRESS": OperatorConfig("replace", {"new_value": "<EMAIL>"}),
        "PHONE_NUMBER": OperatorConfig("mask", {"type": "mask", "masking_char": "*", "chars_to_mask": 4, "from_end": True}),
        "CPF": OperatorConfig("replace", {"new_value": "<CPF/CIN>"}),
        "DATE_TIME": OperatorConfig("keep"),
        "OAB_NUMBER": OperatorConfig("replace", {"new_value": "<OAB>"}),
        "CEP_NUMBER": OperatorConfig("replace", {"new_value": "<CEP>"}),
        "ESTADO_CIVIL": OperatorConfig("keep"),
        "ORGANIZACAO_CONHECIDA": OperatorConfig("keep"),
        "ID_DOCUMENTO": OperatorConfig("keep"),
        "LEGAL_TITLE": OperatorConfig("keep"), # <-- CORREÇÃO 2: REGRA ADICIONADA
        "LEGAL_OR_COMMON_TERM": OperatorConfig("keep"),
        "CNH": OperatorConfig("replace", {"new_value": "***********"}),
        "SIAPE": OperatorConfig("replace", {"new_value": "***"}),
        "RG_NUMBER": OperatorConfig("replace", {"new_value": "<NUMERO RG>"}),
        "MATRICULA_SIAPE": OperatorConfig("replace", {"new_value": "***"}),
        "TERMO_COMUM": OperatorConfig("keep")
    }

def carregar_anonymizer_engine(): return AnonymizerEngine()


# ---> CORREÇÃO 1: INICIALIZAÇÃO CORRIGIDA <---
analyzer_engine = carregar_analyzer_engine(
    LISTA_ESTADOS_CAPITAIS_BR,
    TERMOS_CABECALHO_LEGAL_NAO_ANONIMIZAR,
    LISTA_SOBRENOMES_FREQUENTES_BR,
    LISTA_ESTADO_CIVIL,
    LISTA_ORGANIZACOES_CONHECIDAS,
    LISTA_TERMOS_COMUNS,
    LISTA_TITULOS_LEGAIS
)
anonymizer_engine = carregar_anonymizer_engine() # Corrigido de 'carregar_analyzer_engine'
operadores = obter_operadores_anonimizacao()

def dataframe_entidades_vazio():
    return pd.DataFrame(columns=COLUNAS_ENTIDADES_DETECTADAS)


def gerar_resumo_processamento(df_resultados, origem, tempo_segundos=None):
    sufixo_tempo = ""
    if tempo_segundos is not None:
        sufixo_tempo = f" Tempo total: {tempo_segundos:.2f}s."
    if df_resultados is None or df_resultados.empty:
        return f"**Resumo ({origem}):** nenhuma entidade detectada.{sufixo_tempo}"
    total_entidades = len(df_resultados)
    tipos_unicos = df_resultados["Entidade"].nunique()
    top_tipos = df_resultados["Entidade"].value_counts().head(3)
    top_formatado = ", ".join([f"{entidade} ({qtde})" for entidade, qtde in top_tipos.items()])
    return (
        f"**Resumo ({origem}):** {total_entidades} entidades detectadas em "
        f"{tipos_unicos} tipos. Principais ocorrencias: {top_formatado}.{sufixo_tempo}"
    )

def extrair_texto_de_pdf(caminho_arquivo_pdf):
    texto_completo = ""
    try:
        with fitz.open(caminho_arquivo_pdf) as documento_pdf:
            for pagina in documento_pdf:
                texto_completo += pagina.get_text()
    except Exception as e:
        print(f"Erro ao extrair texto do PDF: {e}")
        return None, f"Erro ao extrair texto do PDF: {e}"
    if not texto_completo.strip():
        return None, "O PDF carregado não contém texto extraível."
    return texto_completo, None


def _limpar_nome_extraido(candidato_nome):
    if not candidato_nome:
        return ""
    nome_limpo = re.sub(r"\s+", " ", candidato_nome).strip()
    nome_limpo = re.sub(r"\s+e\s+outros?\b.*$", "", nome_limpo, flags=re.IGNORECASE)
    nome_limpo = re.sub(r"\s*[-,:;]\s*$", "", nome_limpo).strip()
    return nome_limpo


def _normalizar_texto_para_filtro(texto):
    texto_normalizado = unicodedata.normalize("NFKD", texto)
    texto_sem_acentos = "".join(
        c for c in texto_normalizado if not unicodedata.combining(c)
    )
    return re.sub(r"\s+", " ", texto_sem_acentos).upper().strip()


def _contem_termo_normalizado(texto_normalizado, termo_normalizado):
    padrao = re.compile(
        rf"(?<![A-Z0-9]){re.escape(termo_normalizado)}(?![A-Z0-9])"
    )
    return bool(padrao.search(texto_normalizado))


def _nome_parte_passa_filtros(nome_candidato):
    if not nome_candidato:
        return False
    if "<" in nome_candidato or ">" in nome_candidato:
        return False

    tokens_alfabeticos = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", nome_candidato)
    if len(tokens_alfabeticos) < 2:
        return False

    nome_normalizado = _normalizar_texto_para_filtro(nome_candidato)

    for termo in TERMOS_EXCLUSAO_NOME_PARTE:
        if _contem_termo_normalizado(nome_normalizado, termo):
            return False

    for termo in TERMOS_INDICADORES_PJ:
        if _contem_termo_normalizado(nome_normalizado, termo):
            return False

    return True


def extrair_nomes_parte_alvo(texto_original):
    nomes_encontrados = set()
    if not texto_original or not texto_original.strip():
        return nomes_encontrados

    papeis_regex = "|".join(re.escape(papel) for papel in PAPEIS_ALVO_NOME_PARTE)

    padrao_linha_rotulo = re.compile(
        rf"(?im)^\s*({papeis_regex})\s*:\s*(.+?)\s*$"
    )
    for match in padrao_linha_rotulo.finditer(texto_original):
        nome_limpo = _limpar_nome_extraido(match.group(2))
        if _nome_parte_passa_filtros(nome_limpo):
            nomes_encontrados.add(nome_limpo)

    padrao_linha_partes = re.compile(
        rf"(?im)^\s*(.+?)\s*\(\s*({papeis_regex})\s*\)\s*$"
    )
    for match in padrao_linha_partes.finditer(texto_original):
        nome_limpo = _limpar_nome_extraido(match.group(1))
        if _nome_parte_passa_filtros(nome_limpo):
            nomes_encontrados.add(nome_limpo)

    return nomes_encontrados


def _nome_generico_passa_filtros(nome_candidato):
    if not nome_candidato:
        return False
    if "<" in nome_candidato or ">" in nome_candidato:
        return False

    tokens_alfabeticos = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", nome_candidato)
    if len(tokens_alfabeticos) < 2:
        return False

    nome_normalizado = _normalizar_texto_para_filtro(nome_candidato)

    for termo in TERMOS_INDICADORES_PJ + PALAVRAS_NAO_NOME_GENERICAS:
        if _contem_termo_normalizado(nome_normalizado, termo):
            return False

    return True


def extrair_nomes_pessoais_contextuais(texto_original):
    nomes_encontrados = set()
    if not texto_original or not texto_original.strip():
        return nomes_encontrados

    conectores_nome = r"(?:DE|DA|DO|DAS|DOS|E|de|da|do|das|dos|e)"
    bloco_nome = rf"[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ]+(?:\s+{conectores_nome})?(?:\s+[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ]+){{1,7}}"

    padroes = [
        re.compile(
            rf"(?im)^\s*(?P<nome>{bloco_nome})\s*,\s*(?:brasileir[oa]|solteir[oa]|casad[oa]|desempregad[oa]|portador[oa])\b"
        ),
        re.compile(
            rf"(?im)\b(?:AUTOR(?:A)?|REQUERENTE|RECORRENTE|RECORRIDO(?:A)?|PERIT[AO]|ADVOGAD[OA]|FILIA[ÇC][ÃA]O|ASSINADO\s+ELETRONICAMENTE\s+POR)\s*[:\-]?\s*(?P<nome>{bloco_nome})"
        ),
        re.compile(
            rf"(?im)\(\s*[IVXLC]+\s*\)\s*(?P<nome>{bloco_nome})"
        ),
    ]

    for padrao in padroes:
        for match in padrao.finditer(texto_original):
            nome_limpo = _limpar_nome_extraido(match.group("nome"))
            if _nome_generico_passa_filtros(nome_limpo):
                nomes_encontrados.add(nome_limpo)

    return nomes_encontrados


def anonimizar_nomes_extraidos(texto, nomes_extraidos, placeholder):
    if not texto or not nomes_extraidos:
        return texto
    texto_resultado = texto
    for nome in sorted(nomes_extraidos, key=len, reverse=True):
        if not nome:
            continue
        nome_normalizado = re.sub(r"\s+", " ", nome).strip()
        tokens_nome = [token for token in nome_normalizado.split(" ") if token]
        if not tokens_nome:
            continue
        padrao_nome_flexivel = r"\s+".join(re.escape(token) for token in tokens_nome)
        regex_nome = re.compile(rf"(?i)(?<!\w){padrao_nome_flexivel}(?!\w)")
        texto_resultado = regex_nome.sub(placeholder, texto_resultado)
    return texto_resultado


def normalizar_placeholders_nome_parte(texto):
    if not texto:
        return texto
    return texto.replace(PLACEHOLDER_NOME_PARTE_INTERNO, PLACEHOLDER_NOME_PARTE_EXTERNO)


def _extrair_apenas_digitos(texto):
    return re.sub(r"\D", "", texto or "")


def cpf_tem_digitos_validos(texto_cpf):
    cpf = _extrair_apenas_digitos(texto_cpf)
    if len(cpf) != 11:
        return False
    if cpf == cpf[0] * 11:
        return False

    soma_dv1 = sum(int(cpf[i]) * (10 - i) for i in range(9))
    resto_dv1 = soma_dv1 % 11
    dv1 = 0 if resto_dv1 < 2 else 11 - resto_dv1

    soma_dv2 = sum(int(cpf[i]) * (11 - i) for i in range(10))
    resto_dv2 = soma_dv2 % 11
    dv2 = 0 if resto_dv2 < 2 else 11 - resto_dv2

    return cpf[-2:] == f"{dv1}{dv2}"


def trecho_tem_contexto_cpf(texto_analise, inicio, fim, alcance=32):
    inicio_contexto = max(0, inicio - alcance)
    fim_contexto = min(len(texto_analise), fim + alcance)
    contexto = texto_analise[inicio_contexto:fim_contexto]
    return bool(
        re.search(
            r"(?i)\b(?:cpf|cin|carteira\s+de\s+identidade\s+nacional)\b",
            contexto
        )
    )


def trecho_tem_contexto_oab(texto_analise, inicio, fim, alcance=96):
    inicio_contexto = max(0, inicio - alcance)
    fim_contexto = min(len(texto_analise), fim + alcance)
    contexto = texto_analise[inicio_contexto:fim_contexto]
    return bool(
        re.search(
            r"(?i)\b(?:advogad[oa]|procurador(?:ia)?|representante(?:s)?|polo\s+ativo|polo\s+passivo|oab|assinado\s+eletronicamente|subscr(?:eve|ito)|inscri[çc][aã]o)\b",
            contexto,
        )
    )


def filtrar_resultados_analise(resultados_analise, texto_analise):
    resultados_filtrados = []
    for resultado in resultados_analise:
        trecho_detectado = texto_analise[resultado.start:resultado.end]
        if resultado.entity_type in {"CPF", "CIN"}:
            cpf_valido = cpf_tem_digitos_validos(trecho_detectado)
            cpf_em_contexto = trecho_tem_contexto_cpf(texto_analise, resultado.start, resultado.end)
            if not cpf_valido and not cpf_em_contexto:
                continue
        if resultado.entity_type == "OAB_NUMBER":
            trecho_tem_oab = bool(re.search(r"(?i)\bOAB\b", trecho_detectado))
            if not trecho_tem_oab:
                if not trecho_tem_contexto_oab(texto_analise, resultado.start, resultado.end):
                    continue
        resultados_filtrados.append(resultado)
    return resultados_filtrados


def gerar_arquivo_docx(texto):
    if not texto or texto.strip() == "":
        return None
    try:
        documento = Document()
        documento.add_paragraph(texto)
        filepath = "documento_anonimizado.docx"
        documento.save(filepath)
        return filepath
    except Exception as e:
        gr.Error(f"Erro ao gerar o arquivo .docx: {e}")
        return None


# --- Funções de Lógica da Interface (Event Handlers) ---
def _anonimizar_logica(texto_original):
    nomes_parte_alvo = extrair_nomes_parte_alvo(texto_original)
    texto_para_anonimizar = anonimizar_nomes_extraidos(
        texto_original,
        nomes_parte_alvo,
        PLACEHOLDER_NOME_PARTE_INTERNO,
    )
    nomes_contextuais = extrair_nomes_pessoais_contextuais(texto_para_anonimizar)
    texto_para_anonimizar = anonimizar_nomes_extraidos(
        texto_para_anonimizar,
        nomes_contextuais,
        "<NOME>",
    )
    entidades_para_analise = list(operadores.keys()) + ["SAFE_LOCATION", "LEGAL_HEADER", "ESTADO_CIVIL",
                                                        "ORGANIZACAO_CONHECIDA", "ID_DOCUMENTO", "CNH", "SIAPE", "MATRICULA_SIAPE"]
    entidades_para_analise = list(set(entidades_para_analise) - {"DEFAULT", "PERSON"})
    resultados_analise_brutos = analyzer_engine.analyze(
        text=texto_para_anonimizar, language='pt', entities=entidades_para_analise, return_decision_process=False)
    resultados_analise = filtrar_resultados_analise(resultados_analise_brutos, texto_para_anonimizar)
    resultado_anonimizado_obj = anonymizer_engine.anonymize(
        text=texto_para_anonimizar, analyzer_results=resultados_analise, operators=operadores)
    texto_anonimizado = normalizar_placeholders_nome_parte(resultado_anonimizado_obj.text)
    dados_resultados = [
        {
            "Entidade": res.entity_type,
            "Texto Detectado": normalizar_placeholders_nome_parte(texto_para_anonimizar[res.start:res.end]),
            "Início": res.start,
            "Fim": res.end,
            "Score": f"{res.score:.2f}",
        }
        for res in sorted(resultados_analise, key=lambda x: x.start)
    ]
    return texto_anonimizado, pd.DataFrame(dados_resultados)


def atualizar_estado_botao_texto(texto_original):
    texto_valido = bool(texto_original and texto_original.strip())
    return gr.update(interactive=texto_valido)


def atualizar_estado_botao_pdf(arquivo_temp):
    return gr.update(interactive=arquivo_temp is not None)


def desativar_botao():
    return gr.update(interactive=False)


def processar_texto_area(texto_original):
    if not texto_original or not texto_original.strip():
        gr.Warning("Cole ou digite um texto antes de iniciar a anonimização.")
        return ESTADO_VAZIO_TEXTO_ANONIMIZADO, dataframe_entidades_vazio(), RESUMO_VAZIO_TEXTO
    inicio_processamento = time.perf_counter()
    try:
        texto_anonimizado, df_resultados = _anonimizar_logica(texto_original)
        tempo_total = time.perf_counter() - inicio_processamento
        resumo_processamento = gerar_resumo_processamento(df_resultados, "Texto", tempo_total)
        gr.Info("Texto da área anonimizado com sucesso!")
        return texto_anonimizado, df_resultados, resumo_processamento
    except Exception as e:
        gr.Error(f"Ocorreu um erro durante a anonimização: {e}")
        tempo_total = time.perf_counter() - inicio_processamento
        return "Não foi possível processar o texto.", dataframe_entidades_vazio(), f"**Resumo:** erro no processamento. Tempo total: {tempo_total:.2f}s."


def limpar_texto_area():
    return "", ESTADO_VAZIO_TEXTO_ANONIMIZADO, dataframe_entidades_vazio(), RESUMO_VAZIO_TEXTO, gr.update(interactive=False)


def processar_arquivo_pdf(arquivo_temp, progress=gr.Progress()):
    if arquivo_temp is None:
        gr.Warning("Selecione um arquivo PDF antes de iniciar a anonimização.")
        return ESTADO_VAZIO_PDF_ORIGINAL, ESTADO_VAZIO_PDF_ANONIMIZADO, RESUMO_VAZIO_PDF
    inicio_processamento = time.perf_counter()
    progress(0, desc="Iniciando...")
    try:
        progress(0.2, desc="Extraindo texto do PDF...")
        texto_extraido, erro = extrair_texto_de_pdf(arquivo_temp.name)
        if erro:
            gr.Error(erro)
            tempo_total = time.perf_counter() - inicio_processamento
            return ESTADO_VAZIO_PDF_ORIGINAL, ESTADO_VAZIO_PDF_ANONIMIZADO, f"**Resumo:** erro na extração. Tempo total: {tempo_total:.2f}s."
        progress(0.6, desc="Anonimizando o conteúdo...")
        texto_anonimizado, df_resultados = _anonimizar_logica(texto_extraido)
        tempo_total = time.perf_counter() - inicio_processamento
        resumo_processamento = gerar_resumo_processamento(df_resultados, "PDF", tempo_total)
        progress(1, desc="Concluído!")
        gr.Info("Arquivo PDF anonimizado com sucesso!")
        return texto_extraido, texto_anonimizado, resumo_processamento
    except Exception as e:
        gr.Error(f"Ocorreu um erro ao processar o PDF: {e}")
        tempo_total = time.perf_counter() - inicio_processamento
        return ESTADO_VAZIO_PDF_ORIGINAL, ESTADO_VAZIO_PDF_ANONIMIZADO, f"**Resumo:** erro no processamento. Tempo total: {tempo_total:.2f}s."

# --- CSS Customizado para aumentar o botão de cópia ---
custom_css = """
:root {
    --font-base: "Segoe UI", "Helvetica Neue", Arial, sans-serif;
    --color-primary: #0b5cab;
    --color-primary-hover: #084987;
    --color-bg-subtle: #f5f8fc;
    --color-border: #d8e0eb;
    --space-1: 8px;
    --space-2: 16px;
    --space-3: 24px;
    --radius-1: 10px;
}

.gradio-container {
    font-family: var(--font-base);
    padding: var(--space-3) !important;
}

#header {
    align-items: center;
    gap: var(--space-2);
    margin-bottom: var(--space-2);
}

#header h1 {
    margin-bottom: var(--space-1) !important;
}

.gradio-container .tabs {
    margin-top: var(--space-2);
}

.gradio-container .tabitem {
    padding-top: var(--space-2);
}

.gradio-container .gr-accordion {
    margin-top: var(--space-2);
}

.gradio-container textarea,
.gradio-container input[type="text"],
.gradio-container input[type="file"],
.gradio-container table {
    border-radius: var(--radius-1) !important;
    border-color: var(--color-border) !important;
}

.gradio-container .gr-button {
    border-radius: var(--radius-1) !important;
    padding: var(--space-1) var(--space-2) !important;
}

.gradio-container .gr-button.primary {
    background: var(--color-primary) !important;
    border-color: var(--color-primary) !important;
}

.gradio-container .gr-button.primary:hover {
    background: var(--color-primary-hover) !important;
    border-color: var(--color-primary-hover) !important;
}

.gradio-container .gr-button:not(.primary) {
    background: #f2f4f7 !important;
    color: #1f2937 !important;
    border-color: #cfd8e3 !important;
}

.gradio-container .gr-button:not(.primary):hover {
    background: #e7edf4 !important;
    border-color: #b8c4d1 !important;
}

.gradio-container .cta-row {
    gap: var(--space-1);
    margin-bottom: var(--space-1);
}

.gradio-container .gr-markdown {
    margin-bottom: var(--space-2);
}

.gradio-container > .gr-markdown:nth-of-type(2) {
    background: var(--color-bg-subtle);
    border: 1px solid var(--color-border);
    border-radius: var(--radius-1);
    padding: var(--space-2);
}

button.copy-button {
    padding: 6px !important;
    min-width: 36px !important;
    min-height: 36px !important;
}

button.copy-button > svg {
    transform: scale(1.1);
}
"""

# --- CONSTRUÇÃO DA INTERFACE GRADIO ---
with gr.Blocks(theme=gr.themes.Soft(), title="Anonimizador TRF1", css=custom_css) as demo:

    with gr.Row(elem_id="header"):
        with gr.Column(scale=0, min_width=120):
            gr.Image(value=LOGO_FILE_PATH, interactive=False,
                     show_download_button=False, show_label=False, width=100)
        with gr.Column(scale=4):
            gr.Markdown(
                """
                # Anonimizador SINERGIA
                **Anonimize textos jurídicos com apoio de NLP e regras customizadas.**
                Desenvolvido pelo Projeto Sinergia - TRF1.
                """
            )

    gr.Markdown(
        """
        ### Limites e privacidade
        - Esta ferramenta processa o texto enviado para detectar e anonimizar entidades.
        - PDFs digitalizados como imagem podem não conter texto extraível (OCR não incluso nesta versão).
        - Revise o resultado antes de compartilhar o documento, pois nenhum modelo acerta 100% dos casos.
        - Siga as políticas internas de privacidade e LGPD ao usar dados pessoais.
        """
    )

    with gr.Tabs():
        with gr.TabItem("Anonimizar texto"):
            gr.Markdown("### Passo 1: cole o texto original e clique em **Anonimizar texto**.")
            with gr.Row():
                with gr.Column(scale=1):
                    texto_original_area = gr.Textbox(
                        lines=15,
                        label="Texto original",
                        placeholder="Cole ou digite o conteúdo que deseja anonimizar.")

                with gr.Column(scale=1):
                    texto_anonimizado_area = gr.Textbox(
                        lines=15,
                        value=ESTADO_VAZIO_TEXTO_ANONIMIZADO,
                        label="Texto anonimizado",
                        interactive=False,
                        show_copy_button=True)

            with gr.Row(elem_classes=["cta-row"]):
                btn_anonimizar_area = gr.Button("Anonimizar texto", variant="primary", size="lg", interactive=False)
                btn_limpar_area = gr.Button("Limpar campos", variant="secondary")
            resumo_texto_area = gr.Markdown(value=RESUMO_VAZIO_TEXTO)

            with gr.Accordion("Ver entidades detectadas", open=False):
                resultados_df_area = gr.DataFrame(
                    label="Entidades encontradas",
                    value=dataframe_entidades_vazio(),
                    interactive=False)

        with gr.TabItem("Anonimizar arquivo PDF"):
            gr.Markdown("### Passo 1: envie um PDF com texto pesquisável e clique em **Anonimizar PDF**.")
            upload_pdf = gr.File(label="Selecione o arquivo PDF", file_types=['.pdf'])
            with gr.Row(elem_classes=["cta-row"]):
                btn_anonimizar_pdf = gr.Button("Anonimizar PDF", variant="primary", size="lg", interactive=False)
            resumo_pdf = gr.Markdown(value=RESUMO_VAZIO_PDF)
            gr.Markdown("### Passo 2: revise o texto original extraído e a versão anonimizada.")
            with gr.Row():
                with gr.Accordion("Ver texto original extraído do PDF", open=False):
                    texto_original_pdf = gr.Textbox(
                        lines=15,
                        value=ESTADO_VAZIO_PDF_ORIGINAL,
                        label="Texto original extraído",
                        interactive=False)
                with gr.Column():
                    texto_anonimizado_pdf = gr.Textbox(
                        lines=14,
                        value=ESTADO_VAZIO_PDF_ANONIMIZADO,
                        label="Texto anonimizado",
                        interactive=False,
                        show_copy_button=True)

    # --- Lógica de Conexão dos Componentes (Event Listeners) ---
    texto_original_area.change(fn=atualizar_estado_botao_texto, inputs=[texto_original_area], outputs=[btn_anonimizar_area])
    upload_pdf.change(fn=atualizar_estado_botao_pdf, inputs=[upload_pdf], outputs=[btn_anonimizar_pdf])
    evento_anonimizar_texto = btn_anonimizar_area.click(
        fn=desativar_botao,
        outputs=[btn_anonimizar_area],
        queue=False,
    ).then(
        fn=processar_texto_area,
        inputs=[texto_original_area],
        outputs=[texto_anonimizado_area, resultados_df_area, resumo_texto_area],
    )
    evento_anonimizar_texto.then(
        fn=atualizar_estado_botao_texto,
        inputs=[texto_original_area],
        outputs=[btn_anonimizar_area],
        queue=False,
    )
    btn_limpar_area.click(fn=limpar_texto_area, outputs=[
                          texto_original_area, texto_anonimizado_area, resultados_df_area, resumo_texto_area, btn_anonimizar_area])
    evento_anonimizar_pdf = btn_anonimizar_pdf.click(
        fn=desativar_botao,
        outputs=[btn_anonimizar_pdf],
        queue=False,
    ).then(
        fn=processar_arquivo_pdf,
        inputs=[upload_pdf],
        outputs=[texto_original_pdf, texto_anonimizado_pdf, resumo_pdf],
    )
    evento_anonimizar_pdf.then(
        fn=atualizar_estado_botao_pdf,
        inputs=[upload_pdf],
        outputs=[btn_anonimizar_pdf],
        queue=False,
    )

# --- Ponto de Entrada para Iniciar o App ---
if __name__ == "__main__":
    if analyzer_engine and anonymizer_engine:
        print("Motores de anonimização carregados com sucesso. Iniciando a interface Gradio...")
        demo.launch()
    else:
        print("ERRO CRÍTICO: Não foi possível iniciar a aplicação pois os motores de anonimização falharam ao carregar.")
